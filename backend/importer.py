import csv
import io
import re
import unicodedata
from datetime import date, datetime

from openpyxl import load_workbook

SYNONYMS = {
    "nome": ["nome", "nome vino", "nome del vino", "vino", "etichetta", "wine", "name"],
    "produttore": ["produttore", "produttore / cantina", "produttore o cantina", "cantina", "azienda",
                   "azienda agricola", "producer", "winery"],
    "regione": ["regione", "region", "provenienza", "paese"],
    "denominazione": ["denominazione", "denominazione di origine", "appellation", "doc", "docg"],
    "tipologia": ["tipologia", "tipo", "colore", "categoria", "type", "affinamento", "invecchiamento"],
    "vitigno": ["vitigno", "vitigni", "uvaggio", "uva", "uve", "grape", "grapes"],
    "annata": ["annata", "anno", "vintage", "millesimo"],
    "formato": ["formato", "bottiglia", "capacità", "size"],
    "gradazione": ["gradazione", "gradazione alcolica", "grado alcolico", "alcol", "vol", "% vol", "abv"],
    "temperatura_servizio": ["temperatura di servizio", "temperatura servizio", "temperatura", "servizio",
                             "temp servizio"],
    "prezzo_acquisto": ["prezzo di acquisto", "prezzo acquisto", "acquisto", "costo", "prezzo costo",
                        "prezzo di acquisto (€)", "costo bottiglia"],
    "prezzo_carta": ["prezzo in carta", "prezzo carta", "prezzo in carta (€)", "prezzo", "prezzo vendita",
                     "prezzo di vendita", "carta", "vendita", "price"],
    "preferito": ["preferito", "preferiti", "consigliato"],
    "note": ["note", "note descrittive", "descrizione", "notes", "note di degustazione"],
}

TIPOLOGIE = ["Rosso", "Bianco", "Rosato", "Spumante", "Champagne", "Orange", "Dolce / Passito", "Fortificato"]
TRUE_VALUES = {"si", "sì", "yes", "x", "1", "true", "vero", "*", "y", "s"}
MAX_ROWS = 5000


def norm(value) -> str:
    text = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]", "", text)


LOOKUP = {norm(s): field for field, syns in SYNONYMS.items() for s in syns}


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _price(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    s = re.sub(r"[€\s]|eur", "", v.lower())
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    return round(float(s), 2)


def _annata(v) -> str:
    if isinstance(v, (datetime, date)):
        return str(v.year)
    return _text(v)


def _formato(v) -> str:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return f"{float(v):g}".replace(".", ",") + " l"
    return _text(v)


def _tipologia(v) -> str:
    t = _text(v)
    for known in TIPOLOGIE:
        if t.casefold() == known.casefold():
            return known
    return t[:1].upper() + t[1:]


def _gradazione(v) -> str:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        n = float(v) * 100 if 0 < float(v) < 1 else float(v)
        return f"{n:g}".replace(".", ",") + "% vol"
    return _text(v)


def _temperatura(v) -> str:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return f"{float(v):g}".replace(".", ",") + " °C"
    return _text(v)


CONVERTERS = {"gradazione": _gradazione, "temperatura_servizio": _temperatura,
              "prezzo_acquisto": _price, "prezzo_carta": _price, "annata": _annata, "formato": _formato,
              "tipologia": _tipologia, "preferito": lambda v: _text(v).casefold() in TRUE_VALUES}


def _read_rows(filename: str, raw: bytes) -> list[list[list]]:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception:
            raise ValueError("Impossibile leggere il file Excel")
        sheets = [[list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets]
        wb.close()
        return sheets
    if name.endswith((".csv", ".txt")):
        text = None
        for enc in ("utf-8-sig", "cp1252"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        try:
            delim = csv.Sniffer().sniff(text[:5000], delimiters=";,\t").delimiter
        except csv.Error:
            delim = ";"
        return [list(csv.reader(io.StringIO(text), delimiter=delim))]
    if name.endswith(".xls"):
        raise ValueError("Il formato .xls non è supportato: in Excel usa «Salva con nome» → .xlsx")
    raise ValueError("Formato non supportato: carica un file .xlsx o .csv")


def _find_header(rows):
    for i, row in enumerate(rows[:15]):
        mapping = {}
        for idx, cell in enumerate(row):
            field = LOOKUP.get(norm(cell)) if cell is not None else None
            if field and field not in mapping.values():
                mapping[idx] = field
        if "nome" in mapping.values() and len(mapping) >= 2:
            return i, mapping
    return None


def parse_file(filename: str, raw: bytes) -> dict:
    """Returns {columns: {field: header}, ignored: [headers], rows: [{riga, data}]}."""
    for rows in _read_rows(filename, raw):
        found = _find_header(rows)
        if not found:
            continue
        header_idx, mapping = found
        header = rows[header_idx]
        ignored = [_text(c) for i, c in enumerate(header) if i not in mapping and _text(c)]
        parsed = []
        for offset, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
            if not any(_text(c) for c in row):
                continue
            data, errors = {}, []
            for idx, field in mapping.items():
                value = row[idx] if idx < len(row) else None
                try:
                    data[field] = CONVERTERS.get(field, _text)(value)
                except (ValueError, TypeError):
                    errors.append(f"Valore non valido in «{_text(header[idx])}»: {_text(value)}")
            parsed.append({"riga": offset, "data": data, "errors": errors})
            if len(parsed) > MAX_ROWS:
                raise ValueError(f"Il file supera il limite di {MAX_ROWS} righe")
        return {"columns": {f: _text(header[i]) for i, f in mapping.items()}, "ignored": ignored, "rows": parsed}
    raise ValueError("Non trovo l'intestazione: serve almeno una colonna «Nome» e una tra «Produttore», «Regione», «Prezzo»…")
