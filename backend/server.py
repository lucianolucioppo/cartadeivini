import os
import io
import re
import csv
import json
import base64
import uuid
import asyncio
import logging
from datetime import datetime, timezone, timedelta
from typing import Optional, Literal
from urllib.parse import urlparse

from core import (db, client, logger, now_iso, hash_password, verify_password, create_access_token,
                  create_refresh_token, decode_token, public_user, get_current_user, require_admin,
                  set_auth_cookies, seed_admin, init_storage, put_object, get_object, compress_image, APP_NAME)

from bson import ObjectId
from bson.errors import InvalidId
from fastapi import FastAPI, APIRouter, Request, Response, HTTPException, Depends, UploadFile, File, Form, Query
from fastapi.responses import JSONResponse, StreamingResponse
from starlette.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field, EmailStr, field_validator
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pydantic import ValidationError
from importer import parse_file, TIPOLOGIE

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(name)s - %(levelname)s - %(message)s")

app = FastAPI(title="Carta dei Vini")
api = APIRouter(prefix="/api")

FRONTEND_URL = os.environ["FRONTEND_URL"]


@app.middleware("http")
async def origin_guard(request: Request, call_next):
    if request.method in ("POST", "PUT", "PATCH", "DELETE"):
        origin = request.headers.get("origin")
        if origin:
            host = urlparse(origin).netloc
            allowed = {urlparse(FRONTEND_URL).netloc, request.headers.get("host"),
                       request.headers.get("x-forwarded-host")}
            if host not in allowed:
                return JSONResponse({"detail": "Origine non consentita"}, status_code=403)
    return await call_next(request)


# ---------- models ----------
class LoginIn(BaseModel):
    email: EmailStr
    password: str


class UserCreate(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    email: EmailStr
    password: str = Field(min_length=8, max_length=128)
    role: Literal["staff", "admin"] = "staff"


class PasswordChange(BaseModel):
    current_password: str
    new_password: str = Field(min_length=8, max_length=128)


class WineIn(BaseModel):
    nome: str = Field(min_length=1, max_length=200)
    produttore: str = Field(min_length=1, max_length=200)
    regione: str = Field("", max_length=100)
    denominazione: str = Field("", max_length=200)
    tipologia: str = Field("", max_length=500)
    affinamento: str = Field("", max_length=500)
    vitigno: str = Field("", max_length=300)
    annata: str = Field("", max_length=20)
    formato: str = Field("", max_length=60)
    gradazione: str = Field("", max_length=30)
    temperatura_servizio: str = Field("", max_length=30)
    prezzo_acquisto: Optional[float] = Field(None, ge=0)
    prezzo_carta: Optional[float] = Field(None, ge=0)
    foto_path: Optional[str] = None
    note: str = Field("", max_length=5000)
    preferito: bool = False
    visibile: bool = True

    @field_validator("nome", "produttore", "regione", "denominazione", "tipologia", "affinamento", "vitigno", "annata",
                     "formato", "gradazione", "temperatura_servizio", "note", mode="before")
    @classmethod
    def strip(cls, v):
        return (v or "").strip() if isinstance(v, str) or v is None else str(v).strip()

    @field_validator("nome", "produttore")
    @classmethod
    def not_blank(cls, v):
        if not v:
            raise ValueError("Campo obbligatorio")
        return v


class Wine(WineIn):
    id: str
    created_at: str
    updated_at: str


class FavoriteIn(BaseModel):
    preferito: bool


class SettingsIn(BaseModel):
    carta_pubblica_attiva: bool



# ---------- auth ----------
@api.post("/auth/login")
async def login(body: LoginIn, request: Request, response: Response):
    email = body.email.lower()
    ip = request.headers.get("x-forwarded-for", request.client.host if request.client else "").split(",")[0].strip()
    identifier = f"{ip}:{email}"
    attempt = await db.login_attempts.find_one({"identifier": identifier})
    if attempt and attempt.get("locked_until") and attempt["locked_until"] > now_iso():
        raise HTTPException(status_code=429, detail="Troppi tentativi falliti. Riprova tra 15 minuti.")
    user = await db.users.find_one({"email": email})
    if not user or not verify_password(body.password, user["password_hash"]):
        count = (attempt or {}).get("count", 0) + 1
        update = {"count": count}
        if count >= 5:
            update = {"count": 0, "locked_until": (datetime.now(timezone.utc) + timedelta(minutes=15)).isoformat()}
        await db.login_attempts.update_one({"identifier": identifier}, {"$set": update}, upsert=True)
        raise HTTPException(status_code=401, detail="Email o password non corretti")
    await db.login_attempts.delete_many({"identifier": identifier})
    uid = str(user["_id"])
    set_auth_cookies(response, create_access_token(uid, email), create_refresh_token(uid))
    return public_user(user)


@api.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/", secure=True, samesite="none")
    response.delete_cookie("refresh_token", path="/", secure=True, samesite="none")
    return {"ok": True}


@api.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return user


@api.post("/auth/refresh")
async def refresh(request: Request, response: Response):
    token = request.cookies.get("refresh_token")
    if not token:
        raise HTTPException(status_code=401, detail="Non autenticato")
    payload = decode_token(token, "refresh")
    user = await db.users.find_one({"_id": ObjectId(payload["sub"])})
    if not user:
        raise HTTPException(status_code=401, detail="Utente non trovato")
    set_auth_cookies(response, create_access_token(str(user["_id"]), user["email"]))
    return public_user(user)


@api.post("/auth/password")
async def change_password(body: PasswordChange, user: dict = Depends(get_current_user)):
    doc = await db.users.find_one({"_id": ObjectId(user["id"])})
    if not verify_password(body.current_password, doc["password_hash"]):
        raise HTTPException(status_code=400, detail="La password attuale non è corretta")
    await db.users.update_one({"_id": doc["_id"]}, {"$set": {"password_hash": hash_password(body.new_password)}})
    return {"ok": True}


# ---------- users (admin) ----------
@api.get("/users")
async def list_users(_: dict = Depends(require_admin)):
    users = await db.users.find({}).sort("created_at", 1).to_list(500)
    return [public_user(u) for u in users]


@api.post("/users", status_code=201)
async def create_user(body: UserCreate, _: dict = Depends(require_admin)):
    email = body.email.lower()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Esiste già un utente con questa email")
    doc = {"email": email, "name": body.name.strip(), "role": body.role,
           "password_hash": hash_password(body.password), "created_at": now_iso()}
    res = await db.users.insert_one(doc)
    doc["_id"] = res.inserted_id
    return public_user(doc)


@api.delete("/users/{user_id}")
async def delete_user(user_id: str, admin: dict = Depends(require_admin)):
    if user_id == admin["id"]:
        raise HTTPException(status_code=400, detail="Non puoi eliminare il tuo account")
    try:
        oid = ObjectId(user_id)
    except InvalidId:
        raise HTTPException(status_code=404, detail="Utente non trovato")
    res = await db.users.delete_one({"_id": oid})
    if not res.deleted_count:
        raise HTTPException(status_code=404, detail="Utente non trovato")
    return {"ok": True}


# ---------- wines ----------
SORTS = {"nome": [("nome", 1)], "produttore": [("produttore", 1), ("nome", 1)],
         "annata": [("annata", -1), ("nome", 1)], "prezzo": [("prezzo_carta", 1), ("nome", 1)],
         "prezzo_desc": [("prezzo_carta", -1), ("nome", 1)], "recenti": [("created_at", -1)]}


def wine_filter(q, regione, tipologia, denominazione, annata, produttore, preferito) -> dict:
    flt = {}
    if q:
        rx = {"$regex": re.escape(q.strip()), "$options": "i"}
        flt["$or"] = [{"nome": rx}, {"produttore": rx}]
    for key, val in (("regione", regione), ("tipologia", tipologia), ("denominazione", denominazione),
                     ("annata", annata), ("produttore", produttore)):
        if val:
            flt[key] = val
    if preferito is not None:
        flt["preferito"] = preferito
    return flt


async def fetch_wines(flt: dict, sort: str = "nome") -> list[dict]:
    cursor = db.wines.find(flt, {"_id": 0}).collation({"locale": "it", "strength": 1}).sort(SORTS.get(sort, SORTS["nome"]))
    return [Wine(**w).model_dump() for w in await cursor.to_list(5000)]


@api.get("/wines")
async def list_wines(q: Optional[str] = None, regione: Optional[str] = None, tipologia: Optional[str] = None,
                     denominazione: Optional[str] = None, annata: Optional[str] = None,
                     produttore: Optional[str] = None, preferito: Optional[bool] = None, sort: str = "nome",
                     _: dict = Depends(get_current_user)):
    return await fetch_wines(wine_filter(q, regione, tipologia, denominazione, annata, produttore, preferito), sort)


@api.get("/wines/facets")
async def wine_facets(_: dict = Depends(get_current_user)):
    out = {}
    for key in ("regione", "tipologia", "denominazione", "annata", "produttore", "vitigno", "formato"):
        values = [v for v in await db.wines.distinct(key) if v]
        out[key] = sorted(values, reverse=True) if key == "annata" else sorted(values, key=str.casefold)
    return out


@api.get("/wines/{wine_id}")
async def get_wine(wine_id: str, _: dict = Depends(get_current_user)):
    doc = await db.wines.find_one({"id": wine_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Vino non trovato")
    return Wine(**doc).model_dump()


async def _release_photo(path: Optional[str]):
    if path:
        await db.files.update_one({"storage_path": path}, {"$set": {"is_deleted": True}})


@api.post("/wines", status_code=201)
async def create_wine(body: WineIn, _: dict = Depends(get_current_user)):
    ts = now_iso()
    wine = Wine(**body.model_dump(), id=str(uuid.uuid4()), created_at=ts, updated_at=ts)
    await db.wines.insert_one(wine.model_dump())
    return wine.model_dump()


@api.put("/wines/{wine_id}")
async def update_wine(wine_id: str, body: WineIn, _: dict = Depends(get_current_user)):
    existing = await db.wines.find_one({"id": wine_id}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Vino non trovato")
    data = body.model_dump()
    data["updated_at"] = now_iso()
    await db.wines.update_one({"id": wine_id}, {"$set": data})
    if existing.get("foto_path") and existing["foto_path"] != data["foto_path"]:
        await _release_photo(existing["foto_path"])
    return Wine(**{**existing, **data}).model_dump()


@api.patch("/wines/{wine_id}/preferito")
async def set_favorite(wine_id: str, body: FavoriteIn, _: dict = Depends(get_current_user)):
    res = await db.wines.update_one({"id": wine_id}, {"$set": {"preferito": body.preferito}})
    if not res.matched_count:
        raise HTTPException(status_code=404, detail="Vino non trovato")
    return {"id": wine_id, "preferito": body.preferito}


@api.delete("/wines/{wine_id}")
async def delete_wine(wine_id: str, _: dict = Depends(get_current_user)):
    doc = await db.wines.find_one_and_delete({"id": wine_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Vino non trovato")
    await _release_photo(doc.get("foto_path"))
    return {"ok": True}


# ---------- stats ----------
@api.get("/stats/produttori")
async def producers(_: dict = Depends(get_current_user)):
    pipeline = [
        {"$group": {"_id": "$produttore", "count": {"$sum": 1}, "regioni": {"$addToSet": "$regione"},
                    "tipologie": {"$addToSet": "$tipologia"},
                    "preferiti": {"$sum": {"$cond": ["$preferito", 1, 0]}},
                    "foto": {"$max": "$foto_path"}}},
    ]
    rows = await db.wines.aggregate(pipeline).to_list(5000)
    out = [{"produttore": r["_id"], "count": r["count"], "preferiti": r["preferiti"], "foto_path": r["foto"],
            "regioni": sorted([x for x in r["regioni"] if x]), "tipologie": sorted([x for x in r["tipologie"] if x])}
           for r in rows]
    return sorted(out, key=lambda r: r["produttore"].casefold())


@api.get("/stats/regioni-tipologie")
async def regions_types(_: dict = Depends(get_current_user)):
    rows = await db.wines.aggregate([
        {"$group": {"_id": {"r": "$regione", "t": "$tipologia"}, "count": {"$sum": 1}}}
    ]).to_list(5000)
    regioni, tipologie = {}, {}
    for r in rows:
        reg = r["_id"].get("r") or "Non specificata"
        tip = r["_id"].get("t") or "Non specificata"
        regioni.setdefault(reg, {"nome": reg, "count": 0, "tipologie": {}})
        regioni[reg]["count"] += r["count"]
        regioni[reg]["tipologie"][tip] = regioni[reg]["tipologie"].get(tip, 0) + r["count"]
        tipologie[tip] = tipologie.get(tip, 0) + r["count"]
    total = sum(tipologie.values())
    return {
        "totale": total,
        "regioni": sorted(regioni.values(), key=lambda x: (-x["count"], x["nome"])),
        "tipologie": [{"nome": k, "count": v} for k, v in sorted(tipologie.items(), key=lambda x: (-x[1], x[0]))],
    }


# ---------- settings & public carta ----------
async def read_settings() -> dict:
    doc = await db.settings.find_one({"key": "app"}, {"_id": 0}) or {}
    return {"carta_pubblica_attiva": doc.get("carta_pubblica_attiva", True)}


@api.get("/settings")
async def get_settings(_: dict = Depends(get_current_user)):
    return await read_settings()


@api.put("/settings")
async def update_settings(body: SettingsIn, _: dict = Depends(require_admin)):
    await db.settings.update_one({"key": "app"}, {"$set": body.model_dump()}, upsert=True)
    return await read_settings()


PUBLIC_FIELDS = ["id", "nome", "produttore", "regione", "denominazione", "tipologia", "affinamento", "vitigno",
                 "annata", "formato", "gradazione", "temperatura_servizio", "prezzo_carta", "note", "preferito"]


async def _require_public():
    if not (await read_settings())["carta_pubblica_attiva"]:
        raise HTTPException(status_code=404, detail="La carta digitale non è al momento disponibile")


@api.get("/public/carta")
async def public_carta():
    await _require_public()
    wines = await fetch_wines({"visibile": {"$ne": False}}, "produttore")
    return [{**{k: w.get(k) for k in PUBLIC_FIELDS}, "foto": bool(w.get("foto_path"))} for w in wines]


@api.get("/public/foto/{wine_id}")
async def public_photo(wine_id: str):
    await _require_public()
    w = await db.wines.find_one({"id": wine_id, "visibile": {"$ne": False}}, {"_id": 0, "foto_path": 1})
    if not w or not w.get("foto_path"):
        raise HTTPException(status_code=404, detail="Foto non disponibile")
    data, ctype = await load_object(w["foto_path"], "Foto non disponibile")
    return Response(content=data, media_type=ctype, headers={"Cache-Control": "public, max-age=3600"})


# ---------- photos ----------
MAX_UPLOAD = 15 * 1024 * 1024


async def load_object(path: str, detail: str) -> tuple[bytes, str]:
    try:
        return await get_object(path)
    except Exception:
        raise HTTPException(status_code=404, detail=detail)


async def _compress_upload(raw: bytes) -> bytes:
    try:
        return await asyncio.to_thread(compress_image, raw)
    except Exception:
        raise HTTPException(status_code=400, detail="Immagine non leggibile")


async def _store_photo(path: str, data: bytes) -> dict:
    try:
        return await put_object(path, data, "image/jpeg")
    except Exception as e:
        logger.error(f"Upload failed: {e}")
        raise HTTPException(status_code=502, detail="Caricamento della foto non riuscito, riprova")


@api.post("/upload")
async def upload_photo(file: UploadFile = File(...), _: dict = Depends(get_current_user)):
    if not (file.content_type or "").startswith("image/"):
        raise HTTPException(status_code=400, detail="Il file deve essere un'immagine")
    raw = await file.read()
    if len(raw) > MAX_UPLOAD:
        raise HTTPException(status_code=400, detail="Immagine troppo grande (max 15 MB)")
    data = await _compress_upload(raw)
    path = f"{APP_NAME}/etichette/{uuid.uuid4()}.jpg"
    result = await _store_photo(path, data)
    await db.files.insert_one({"id": str(uuid.uuid4()), "storage_path": result["path"], "content_type": "image/jpeg",
                               "original_filename": file.filename, "size": result.get("size", len(data)),
                               "is_deleted": False, "created_at": now_iso()})
    return {"path": result["path"]}


@api.get("/files/{path:path}")
async def get_file(path: str, _: dict = Depends(get_current_user)):
    record = await db.files.find_one({"storage_path": path}, {"_id": 0})
    if not record:
        raise HTTPException(status_code=404, detail="File non trovato")
    data, ctype = await load_object(path, "File non disponibile")
    return Response(content=data, media_type=record.get("content_type") or ctype,
                    headers={"Cache-Control": "private, max-age=604800, immutable"})


# ---------- export ----------
COLUMNS = [("nome", "Nome"), ("produttore", "Produttore / Cantina"), ("regione", "Regione"),
           ("denominazione", "Denominazione"), ("tipologia", "Tipologia"), ("vitigno", "Vitigno"),
           ("affinamento", "Affinamento"),
           ("annata", "Annata"), ("formato", "Formato"), ("gradazione", "Gradazione alcolica"),
           ("temperatura_servizio", "Temperatura di servizio"), ("prezzo_acquisto", "Prezzo di acquisto (€)"),
           ("prezzo_carta", "Prezzo in carta (€)"), ("preferito", "Preferito"),
           ("visibile", "Visibile in carta digitale"), ("note", "Note")]


def _stamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d")


def _cell(key, val, for_csv):
    if key in ("preferito", "visibile"):
        return "Sì" if val else "No"
    if key.startswith("prezzo"):
        if val is None:
            return ""
        return f"{val:.2f}".replace(".", ",") if for_csv else float(val)
    return val or ""


@api.get("/export/csv")
async def export_csv(_: dict = Depends(get_current_user)):
    wines = await fetch_wines({}, "produttore")
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    writer.writerow([label for _, label in COLUMNS])
    for w in wines:
        writer.writerow([_cell(k, w.get(k), True) for k, _ in COLUMNS])
    content = "\ufeff" + buf.getvalue()
    return Response(content=content.encode("utf-8"), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": f'attachment; filename="carta-dei-vini-{_stamp()}.csv"'})


def build_xlsx(wines: list[dict]) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Carta dei Vini"
    ws.append([label for _, label in COLUMNS])
    header_fill = PatternFill("solid", fgColor="5A1220")
    for c in ws[1]:
        c.font = Font(bold=True, color="F3E3B3")
        c.fill = header_fill
        c.alignment = Alignment(vertical="center")
    for w in wines:
        ws.append([_cell(k, w.get(k), False) for k, _ in COLUMNS])
    widths = [30, 28, 16, 28, 14, 26, 40, 10, 14, 18, 22, 18, 18, 10, 14, 60]
    for i, width in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = width
    price_col = [k for k, _ in COLUMNS].index("prezzo_acquisto") + 1
    for row in ws.iter_rows(min_row=2, min_col=price_col, max_col=price_col + 1):
        for c in row:
            c.number_format = '#,##0.00 "€"'
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@api.get("/export/xlsx")
async def export_xlsx(_: dict = Depends(get_current_user)):
    out = build_xlsx(await fetch_wines({}, "produttore"))
    return StreamingResponse(out, media_type=XLSX_TYPE,
                             headers={"Content-Disposition": f'attachment; filename="carta-dei-vini-{_stamp()}.xlsx"'})


# ---------- import ----------
MAX_IMPORT = 10 * 1024 * 1024
FIELD_LABELS = {"nome": "Nome", "produttore": "Produttore", "prezzo_acquisto": "Prezzo di acquisto",
                "prezzo_carta": "Prezzo in carta"}


@api.get("/import/template")
async def import_template(_: dict = Depends(get_current_user)):
    return StreamingResponse(build_xlsx([]), media_type=XLSX_TYPE,
                             headers={"Content-Disposition": 'attachment; filename="modello-carta-dei-vini.xlsx"'})


def _validation_messages(err: ValidationError) -> list[str]:
    out = []
    for e in err.errors():
        field = e["loc"][0] if e["loc"] else ""
        label = FIELD_LABELS.get(field, str(field).capitalize())
        out.append(f"{label} mancante" if field in ("nome", "produttore") else f"{label} non valido")
    return out


async def _read_import(file: UploadFile) -> dict:
    raw = await file.read()
    if len(raw) > MAX_IMPORT:
        raise HTTPException(status_code=400, detail="File troppo grande (max 10 MB)")
    try:
        return await asyncio.to_thread(parse_file, file.filename, raw)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


def _import_key(w: dict, with_formato: bool) -> tuple:
    base = (w["nome"].casefold(), w["produttore"].casefold(), (w.get("annata") or "").casefold())
    return base + ((w.get("formato") or "").casefold(),) if with_formato else base


def _validate_row(r: dict, fields: set) -> tuple[Optional[dict], list[str]]:
    if r["errors"]:
        return None, list(r["errors"])
    try:
        return WineIn(**r["data"]).model_dump(include=fields), []
    except ValidationError as e:
        return None, _validation_messages(e)


def _classify_row(r: dict, fields: set, with_formato: bool, existing: dict, seen: set) -> dict:
    wine, errors = _validate_row(r, fields)
    stato, existing_id = "errore", None
    if wine:
        k = _import_key(wine, with_formato)
        if k in seen:
            errors = ["Riga ripetuta nel file"]
        elif k in existing:
            stato, existing_id = "presente", existing[k]
        else:
            stato = "nuovo"
        seen.add(k)
    d = r["data"]
    return {"riga": r["riga"], "stato": stato, "errori": errors, "wine": wine, "existing_id": existing_id,
            "nome": d.get("nome", ""), "produttore": d.get("produttore", ""), "annata": d.get("annata", ""),
            "prezzo_carta": d.get("prezzo_carta")}


async def analyze_import(file: UploadFile) -> dict:
    parsed = await _read_import(file)
    fields = set(parsed["columns"])
    with_formato = "formato" in fields
    docs = await db.wines.find({}, {"_id": 0, "id": 1, "nome": 1, "produttore": 1, "annata": 1, "formato": 1}).to_list(10000)
    existing = {_import_key(w, with_formato): w["id"] for w in docs}
    seen: set = set()
    rows = [_classify_row(r, fields, with_formato, existing, seen) for r in parsed["rows"]]
    count = lambda s: sum(1 for x in rows if x["stato"] == s)
    return {"columns": parsed["columns"], "ignored": parsed["ignored"], "rows": rows,
            "totali": {"righe": len(rows), "nuovi": count("nuovo"), "presenti": count("presente"), "errori": count("errore")}}


@api.post("/import/preview")
async def import_preview(file: UploadFile = File(...), _: dict = Depends(get_current_user)):
    result = await analyze_import(file)
    preview = [{k: v for k, v in r.items() if k not in ("wine", "existing_id")} for r in result["rows"][:300]]
    return {**result, "rows": preview}


@api.post("/import")
async def import_wines(file: UploadFile = File(...), mode: Literal["skip", "update"] = Form("skip"),
                       _: dict = Depends(get_current_user)):
    result = await analyze_import(file)
    ts = now_iso()
    new_docs, updated = [], 0
    for r in result["rows"]:
        if r["stato"] == "nuovo":
            base = WineIn(**r["wine"]).model_dump()
            new_docs.append(Wine(**base, id=str(uuid.uuid4()), created_at=ts, updated_at=ts).model_dump())
        elif r["stato"] == "presente" and mode == "update":
            await db.wines.update_one({"id": r["existing_id"]}, {"$set": {**r["wine"], "updated_at": ts}})
            updated += 1
    if new_docs:
        await db.wines.insert_many(new_docs)
    t = result["totali"]
    return {"inseriti": len(new_docs), "aggiornati": updated,
            "saltati": t["presenti"] - updated, "errori": t["errori"]}


# ---------- backup / restore ----------
@api.get("/backup")
async def backup(_: dict = Depends(get_current_user)):
    wines = await fetch_wines({}, "nome")
    photos, missing = {}, 0
    for w in wines:
        path = w.get("foto_path")
        if path and path not in photos:
            try:
                data, ctype = await get_object(path)
                photos[path] = {"content_type": ctype, "data": base64.b64encode(data).decode()}
            except Exception:
                missing += 1
    payload = {"app": "carta-dei-vini", "version": 1, "exported_at": now_iso(), "wines": wines,
               "photos": photos, "missing_photos": missing}
    return Response(content=json.dumps(payload, ensure_ascii=False), media_type="application/json",
                    headers={"Content-Disposition": f'attachment; filename="backup-carta-dei-vini-{_stamp()}.json"'})


def _backup_wine(raw: dict) -> Wine:
    ts = now_iso()
    try:
        return Wine(**WineIn(**raw).model_dump(), id=raw.get("id") or str(uuid.uuid4()),
                    created_at=raw.get("created_at") or ts, updated_at=raw.get("updated_at") or ts)
    except ValidationError:
        raise HTTPException(status_code=400, detail=f"Vino non valido nel backup: {raw.get('nome', '?')}")


def _load_json(raw: bytes):
    try:
        return json.loads(raw)
    except ValueError:
        raise HTTPException(status_code=400, detail="File di backup non valido")


def _parse_backup(raw: bytes) -> tuple[list[Wine], dict]:
    payload = _load_json(raw)
    if not isinstance(payload, dict) or payload.get("app") != "carta-dei-vini" or not isinstance(payload.get("wines"), list):
        raise HTTPException(status_code=400, detail="Il file non è un backup di Carta dei Vini")
    raw_photos = payload.get("photos") or {}
    if not isinstance(raw_photos, dict) or not all(isinstance(w, dict) for w in payload["wines"]):
        raise HTTPException(status_code=400, detail="File di backup non valido")
    photos = {p: i for p, i in raw_photos.items() if p.startswith(f"{APP_NAME}/") and isinstance(i, dict)}
    return [_backup_wine(w) for w in payload["wines"]], photos


async def _restore_photo(path: str, info: dict):
    data = base64.b64decode(info["data"])
    ctype = info.get("content_type") or "image/jpeg"
    result = await put_object(path, data, ctype)
    await db.files.update_one({"storage_path": result["path"]},
                              {"$set": {"storage_path": result["path"], "content_type": ctype, "is_deleted": False, "size": len(data)},
                               "$setOnInsert": {"id": str(uuid.uuid4()), "created_at": now_iso()}}, upsert=True)


@api.post("/backup/restore")
async def restore(file: UploadFile = File(...), _: dict = Depends(require_admin)):
    wines, photos = _parse_backup(await file.read())
    for path, info in photos.items():
        await _restore_photo(path, info)
    for w in wines:
        if w.foto_path and w.foto_path not in photos and not await db.files.find_one({"storage_path": w.foto_path}):
            w.foto_path = None
    await db.wines.delete_many({})
    if wines:
        await db.wines.insert_many([w.model_dump() for w in wines])
    return {"vini": len(wines), "foto": len(photos)}


app.include_router(api)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=[FRONTEND_URL, "http://localhost:3000"],
    allow_methods=["*"],
    allow_headers=["*"],
)


async def migrate_affinamento():
    known = {t.casefold() for t in TIPOLOGIE}
    async for w in db.wines.find({"affinamento": {"$exists": False}}, {"_id": 0, "id": 1, "tipologia": 1}):
        t = w.get("tipologia") or ""
        moved = bool(t) and t.casefold() not in known
        update = {"affinamento": t, "tipologia": ""} if moved else {"affinamento": ""}
        await db.wines.update_one({"id": w["id"]}, {"$set": update})


@app.on_event("startup")
async def startup():
    await db.users.create_index("email", unique=True)
    await db.login_attempts.create_index("identifier")
    await db.wines.create_index("id", unique=True)
    await db.files.create_index("storage_path")
    await seed_admin()
    await migrate_affinamento()
    try:
        await asyncio.to_thread(init_storage)
        logger.info("Storage initialized")
    except Exception as e:
        logger.error(f"Storage init failed: {e}")


@app.on_event("shutdown")
async def shutdown():
    client.close()
