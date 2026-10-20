"""Iteration 10: restaurant name setting (/api/settings, /api/public/info) and 'esaurito' flag on wines."""
import csv
import io
import os
import uuid
from pathlib import Path

import pytest
import requests
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook

BASE_URL = os.environ["REACT_APP_BACKEND_URL"].rstrip("/")
API = f"{BASE_URL}/api"
load_dotenv(Path(__file__).resolve().parents[1] / ".env")
ADMIN_EMAIL = os.environ["ADMIN_EMAIL"]
ADMIN_PASSWORD = os.environ["ADMIN_PASSWORD"]
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
STAFF_EMAIL = f"test_staff_{uuid.uuid4().hex[:6]}@example.com"


@pytest.fixture(scope="module")
def admin():
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD}, timeout=30)
    assert r.status_code == 200, r.text
    return s


@pytest.fixture(scope="module")
def staff(admin):
    r = admin.post(f"{API}/users", json={"name": "TEST_Staff", "email": STAFF_EMAIL, "password": "Staffpass-123",
                                         "role": "staff"}, timeout=30)
    assert r.status_code == 201, r.text
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json={"email": STAFF_EMAIL, "password": "Staffpass-123"}, timeout=30)
    assert r.status_code == 200, r.text
    yield s
    for u in admin.get(f"{API}/users", timeout=30).json():
        if u["email"] == STAFF_EMAIL:
            admin.delete(f"{API}/users/{u['id']}", timeout=30)


@pytest.fixture(scope="module", autouse=True)
def _cleanup(admin):
    yield
    admin.put(f"{API}/settings", json={"nome_ristorante": "", "carta_pubblica_attiva": True}, timeout=30)
    for w in admin.get(f"{API}/wines", timeout=30).json():
        if w["nome"].startswith("TEST_"):
            admin.delete(f"{API}/wines/{w['id']}", timeout=30)


# ---------- restaurant name ----------
class TestRestaurantName:
    def test_public_info_no_auth(self):
        r = requests.get(f"{API}/public/info", timeout=30)
        assert r.status_code == 200
        assert set(r.json().keys()) == {"nome_ristorante"}

    def test_set_name_trimmed_and_persisted(self, admin):
        r = admin.put(f"{API}/settings", json={"nome_ristorante": "  TEST_Trattoria  "}, timeout=30)
        assert r.status_code == 200, r.text
        assert r.json()["nome_ristorante"] == "TEST_Trattoria"
        assert r.json()["carta_pubblica_attiva"] is True
        assert admin.get(f"{API}/settings", timeout=30).json()["nome_ristorante"] == "TEST_Trattoria"
        assert requests.get(f"{API}/public/info", timeout=30).json()["nome_ristorante"] == "TEST_Trattoria"

    def test_toggle_carta_does_not_change_name(self, admin):
        r = admin.put(f"{API}/settings", json={"carta_pubblica_attiva": False}, timeout=30)
        assert r.json() == {"carta_pubblica_attiva": False, "nome_ristorante": "TEST_Trattoria"}
        assert requests.get(f"{API}/public/carta", timeout=30).status_code == 404
        r = admin.put(f"{API}/settings", json={"nome_ristorante": "TEST_Due"}, timeout=30)
        assert r.json() == {"carta_pubblica_attiva": False, "nome_ristorante": "TEST_Due"}
        r = admin.put(f"{API}/settings", json={"carta_pubblica_attiva": True}, timeout=30)
        assert r.json() == {"carta_pubblica_attiva": True, "nome_ristorante": "TEST_Due"}

    def test_max_length(self, admin):
        assert admin.put(f"{API}/settings", json={"nome_ristorante": "x" * 121}, timeout=30).status_code == 422
        r = admin.put(f"{API}/settings", json={"nome_ristorante": "x" * 120}, timeout=30)
        assert r.status_code == 200 and len(r.json()["nome_ristorante"]) == 120

    def test_empty_and_whitespace(self, admin):
        r = admin.put(f"{API}/settings", json={"nome_ristorante": "   "}, timeout=30)
        assert r.status_code == 200 and r.json()["nome_ristorante"] == ""

    def test_staff_forbidden_but_can_read(self, staff):
        assert staff.put(f"{API}/settings", json={"nome_ristorante": "hack"}, timeout=30).status_code == 403
        r = staff.get(f"{API}/settings", timeout=30)
        assert r.status_code == 200 and "nome_ristorante" in r.json()

    def test_unauth_put_401(self):
        assert requests.put(f"{API}/settings", json={"nome_ristorante": "x"}, timeout=30).status_code == 401


# ---------- esaurito ----------
def _wine(nome, **kw):
    return {"nome": nome, "produttore": "TEST_Cantina10", "tipologia": "Rosso", "prezzo_carta": 30,
            "prezzo_acquisto": 12, "fornitore": "TEST_Forn", **kw}


class TestEsaurito:
    def test_default_false(self, admin):
        r = admin.post(f"{API}/wines", json=_wine("TEST_Default10"), timeout=30)
        assert r.status_code == 201 and r.json()["esaurito"] is False

    def test_create_update_esaurito(self, admin):
        r = admin.post(f"{API}/wines", json=_wine("TEST_Sold10", esaurito=True), timeout=30)
        assert r.status_code == 201 and r.json()["esaurito"] is True
        wid = r.json()["id"]
        assert admin.get(f"{API}/wines/{wid}", timeout=30).json()["esaurito"] is True
        lst = {w["id"]: w for w in admin.get(f"{API}/wines", timeout=30).json()}
        assert lst[wid]["esaurito"] is True
        body = {**_wine("TEST_Sold10"), "esaurito": False}
        assert admin.put(f"{API}/wines/{wid}", json=body, timeout=30).json()["esaurito"] is False
        assert admin.get(f"{API}/wines/{wid}", timeout=30).json()["esaurito"] is False
        admin.put(f"{API}/wines/{wid}", json={**body, "esaurito": True}, timeout=30)

    def test_public_carta_has_flag_no_private_fields(self, admin):
        r = requests.get(f"{API}/public/carta", timeout=30)
        assert r.status_code == 200
        items = r.json()
        sold = [w for w in items if w["nome"] == "TEST_Sold10"]
        assert sold and sold[0]["esaurito"] is True
        for w in items:
            assert "prezzo_acquisto" not in w and "fornitore" not in w and "foto_path" not in w

    def test_export_csv_column(self, admin):
        r = admin.get(f"{API}/export/csv", timeout=30)
        assert r.status_code == 200
        rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))
        idx = rows[0].index("Esaurito")
        name_idx = rows[0].index("Nome")
        vals = {row[name_idx]: row[idx] for row in rows[1:]}
        assert vals["TEST_Sold10"] == "Sì" and vals["TEST_Default10"] == "No"

    def test_export_xlsx_column(self, admin):
        r = admin.get(f"{API}/export/xlsx", timeout=30)
        assert r.status_code == 200
        ws = load_workbook(io.BytesIO(r.content)).active
        rows = list(ws.iter_rows(values_only=True))
        idx = list(rows[0]).index("Esaurito")
        vals = {row[0]: row[idx] for row in rows[1:]}
        assert vals["TEST_Sold10"] == "Sì" and vals["TEST_Default10"] == "No"
        assert list(rows[0])[-2:] == ("Note", "Abbinamenti") or list(rows[0])[-2:] == ["Note", "Abbinamenti"]  # iter12: Abbinamenti added after Note

    def test_import_esaurito_mapping(self, admin):
        wb = Workbook()
        ws = wb.active
        ws.append(["Nome", "Produttore", "Prezzo in carta", "Esaurito"])
        ws.append(["TEST_ImpSi10", "TEST_Cantina10", 20, "sì"])
        ws.append(["TEST_ImpNo10", "TEST_Cantina10", 21, ""])
        ws.append(["TEST_ImpX10", "TEST_Cantina10", 22, "x"])
        buf = io.BytesIO()
        wb.save(buf)
        r = admin.post(f"{API}/import", files={"file": ("e.xlsx", buf.getvalue(), XLSX_MIME)},
                       data={"mode": "skip"}, timeout=30)
        assert r.status_code == 200, r.text
        assert r.json()["inseriti"] == 3
        w = {x["nome"]: x for x in admin.get(f"{API}/wines", timeout=30).json()}
        assert w["TEST_ImpSi10"]["esaurito"] is True
        assert w["TEST_ImpNo10"]["esaurito"] is False
        assert w["TEST_ImpX10"]["esaurito"] is True

    def test_import_update_without_column_keeps_esaurito(self, admin):
        wb = Workbook()
        ws = wb.active
        ws.append(["Nome", "Produttore", "Prezzo in carta"])
        ws.append(["TEST_ImpSi10", "TEST_Cantina10", 25])
        buf = io.BytesIO()
        wb.save(buf)
        r = admin.post(f"{API}/import", files={"file": ("e.xlsx", buf.getvalue(), XLSX_MIME)},
                       data={"mode": "update"}, timeout=30)
        assert r.json()["aggiornati"] == 1
        w = {x["nome"]: x for x in admin.get(f"{API}/wines", timeout=30).json()}
        assert w["TEST_ImpSi10"]["esaurito"] is True and w["TEST_ImpSi10"]["prezzo_carta"] == 25

    def test_csv_roundtrip_esaurito(self, admin):
        r = admin.get(f"{API}/export/csv", timeout=30)
        r = admin.post(f"{API}/import/preview", files={"file": ("x.csv", r.content, "text/csv")}, timeout=30)
        assert r.status_code == 200
        assert "esaurito" in r.json()["columns"]
        assert r.json()["totali"]["errori"] == 0 and r.json()["totali"]["nuovi"] == 0
