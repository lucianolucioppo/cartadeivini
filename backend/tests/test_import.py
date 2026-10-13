"""Backend tests for Carta dei Vini: Excel/CSV import feature."""
import io
import os
import csv
import pytest
import requests
from openpyxl import Workbook, load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://carta-dei-vini.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"

ADMIN_EMAIL = "admin@cartadeivini.it"
ADMIN_PASSWORD = "Bordeaux-2026!"

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---------- helpers ----------
def _xlsx(rows, before_header_rows=0):
    wb = Workbook()
    ws = wb.active
    for _ in range(before_header_rows):
        ws.append(["", "", ""])  # non-matching
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _csv(rows, delim=";"):
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=delim)
    for r in rows:
        w.writerow(r)
    return buf.getvalue().encode("utf-8")


@pytest.fixture(scope="module")
def admin():
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD}, timeout=30)
    assert r.status_code == 200, f"Admin login failed: {r.text}"
    return s


@pytest.fixture(scope="module", autouse=True)
def _cleanup(admin):
    yield
    # delete anything we created
    r = admin.get(f"{API}/wines", timeout=30)
    for w in r.json():
        if w["nome"].startswith("TEST_"):
            admin.delete(f"{API}/wines/{w['id']}", timeout=30)


# ---------- auth guard ----------
class TestAuth:
    def test_preview_requires_auth(self):
        data = _xlsx([["Nome", "Produttore"], ["TEST_X", "TEST_Y"]])
        r = requests.post(f"{API}/import/preview", files={"file": ("x.xlsx", data, XLSX_MIME)}, timeout=30)
        assert r.status_code == 401

    def test_import_requires_auth(self):
        data = _xlsx([["Nome", "Produttore"], ["TEST_X", "TEST_Y"]])
        r = requests.post(f"{API}/import", files={"file": ("x.xlsx", data, XLSX_MIME)}, data={"mode": "skip"}, timeout=30)
        assert r.status_code == 401

    def test_template_requires_auth(self):
        r = requests.get(f"{API}/import/template", timeout=30)
        assert r.status_code == 401


# ---------- preview ----------
class TestPreview:
    def test_preview_basic_columns_and_synonyms(self, admin):
        data = _xlsx([
            ["Vino", "Cantina", "Tipo", "Anno", "Prezzo", "Colonna Ignota"],
            ["TEST_Barolo", "TEST_Produttore A", "rosso", 2018, "32,50 €", "ciao"],
            ["TEST_Franciacorta", "TEST_Produttore B", "spumante", 2020, 28, "x"],
        ])
        r = admin.post(f"{API}/import/preview", files={"file": ("c.xlsx", data, XLSX_MIME)}, timeout=30)
        assert r.status_code == 200, r.text
        j = r.json()
        cols = j["columns"]
        # synonyms mapped
        assert cols.get("nome") == "Vino"
        assert cols.get("produttore") == "Cantina"
        assert cols.get("tipologia") == "Tipo"
        assert cols.get("annata") == "Anno"
        assert cols.get("prezzo_carta") == "Prezzo"
        # ignored column surfaced
        assert "Colonna Ignota" in j["ignored"]
        # totals: 2 new rows, no existing (fresh DB), no errors
        t = j["totali"]
        assert t["righe"] == 2
        assert t["nuovi"] == 2
        assert t["errori"] == 0
        # row details
        row0 = j["rows"][0]
        assert row0["stato"] == "nuovo"
        assert row0["nome"] == "TEST_Barolo"
        # 32,50 € parsed to 32.5
        assert abs(row0["prezzo_carta"] - 32.5) < 0.001

    def test_preview_tipologia_normalized_and_header_not_first_row(self, admin):
        data = _xlsx([
            ["Nome", "Produttore", "Tipologia"],
            ["TEST_W1", "TEST_P1", "rosso"],
        ], before_header_rows=2)
        r = admin.post(f"{API}/import/preview", files={"file": ("c.xlsx", data, XLSX_MIME)}, timeout=30)
        assert r.status_code == 200, r.text
        j = r.json()
        assert j["totali"]["righe"] == 1
        # Normalization check requires fetching wine data; use import then list
        # Just check preview captured raw values; we test normalization via actual import below.
        assert j["rows"][0]["nome"] == "TEST_W1"

    def test_preview_errors_missing_name_and_invalid_price(self, admin):
        data = _xlsx([
            ["Nome", "Produttore", "Prezzo in carta"],
            ["", "TEST_P", "20,00"],  # missing name
            ["TEST_W2", "TEST_P", "non-numero"],  # invalid price
        ])
        r = admin.post(f"{API}/import/preview", files={"file": ("c.xlsx", data, XLSX_MIME)}, timeout=30)
        assert r.status_code == 200, r.text
        j = r.json()
        errs = [row for row in j["rows"] if row["stato"] == "errore"]
        assert len(errs) == 2
        joined = " ".join(e["errori"][0] for e in errs).lower()
        assert "nome" in joined
        assert "prezzo" in joined

    def test_preview_duplicate_rows_in_file(self, admin):
        data = _xlsx([
            ["Nome", "Produttore", "Annata"],
            ["TEST_Dup", "TEST_PD", 2019],
            ["TEST_Dup", "TEST_PD", 2019],
        ])
        r = admin.post(f"{API}/import/preview", files={"file": ("c.xlsx", data, XLSX_MIME)}, timeout=30)
        assert r.status_code == 200, r.text
        j = r.json()
        assert j["totali"]["nuovi"] == 1
        assert j["totali"]["errori"] == 1
        err_row = [x for x in j["rows"] if x["stato"] == "errore"][0]
        assert "ripetuta" in err_row["errori"][0].lower()

    def test_preview_xls_rejected_italian_message(self, admin):
        r = admin.post(f"{API}/import/preview",
                       files={"file": ("vecchio.xls", b"\xd0\xcf\x11\xe0fake", "application/vnd.ms-excel")}, timeout=30)
        assert r.status_code == 400
        detail = r.json().get("detail", "")
        assert ".xls" in detail.lower()
        assert "xlsx" in detail.lower()

    def test_preview_no_header_rejected(self, admin):
        data = _xlsx([
            ["foo", "bar", "baz"],
            ["abc", "def", "ghi"],
        ])
        r = admin.post(f"{API}/import/preview", files={"file": ("x.xlsx", data, XLSX_MIME)}, timeout=30)
        assert r.status_code == 400
        assert "intestazione" in r.json()["detail"].lower()


# ---------- import ----------
class TestImport:
    def test_import_insert_new(self, admin):
        data = _xlsx([
            ["Nome", "Produttore", "Annata", "Prezzo in carta", "Tipologia"],
            ["TEST_Imp1", "TEST_Prod1", 2019, "25,00", "rosso"],
            ["TEST_Imp2", "TEST_Prod1", 2020, 30, "bianco"],
        ])
        r = admin.post(f"{API}/import", files={"file": ("c.xlsx", data, XLSX_MIME)},
                       data={"mode": "skip"}, timeout=30)
        assert r.status_code == 200, r.text
        j = r.json()
        assert j["inseriti"] == 2
        assert j["aggiornati"] == 0
        # Verify persisted and tipologia normalized
        r = admin.get(f"{API}/wines", timeout=30)
        wines = {w["nome"]: w for w in r.json()}
        assert "TEST_Imp1" in wines
        assert wines["TEST_Imp1"]["tipologia"] == "Rosso"
        assert wines["TEST_Imp1"]["prezzo_carta"] == 25.0
        assert wines["TEST_Imp2"]["tipologia"] == "Bianco"

    def test_import_mode_skip_leaves_existing(self, admin):
        # TEST_Imp1 already exists from prior test
        data = _xlsx([
            ["Nome", "Produttore", "Annata", "Prezzo in carta"],
            ["TEST_Imp1", "TEST_Prod1", 2019, "99,00"],
            ["TEST_Imp3", "TEST_Prod1", 2021, "40,00"],
        ])
        r = admin.post(f"{API}/import", files={"file": ("c.xlsx", data, XLSX_MIME)},
                       data={"mode": "skip"}, timeout=30)
        assert r.status_code == 200
        j = r.json()
        assert j["inseriti"] == 1
        assert j["aggiornati"] == 0
        assert j["saltati"] == 1
        # Price unchanged
        r = admin.get(f"{API}/wines", timeout=30)
        w = next(x for x in r.json() if x["nome"] == "TEST_Imp1")
        assert w["prezzo_carta"] == 25.0

    def test_import_mode_update(self, admin):
        data = _xlsx([
            ["Nome", "Produttore", "Annata", "Prezzo in carta"],
            ["TEST_Imp1", "TEST_Prod1", 2019, "77,50"],
        ])
        r = admin.post(f"{API}/import", files={"file": ("c.xlsx", data, XLSX_MIME)},
                       data={"mode": "update"}, timeout=30)
        assert r.status_code == 200
        j = r.json()
        assert j["aggiornati"] == 1
        r = admin.get(f"{API}/wines", timeout=30)
        w = next(x for x in r.json() if x["nome"] == "TEST_Imp1")
        assert w["prezzo_carta"] == 77.5

    def test_import_update_keeps_foto_path(self, admin):
        # create a wine with a foto_path, then update via import; foto_path must stay
        r = admin.post(f"{API}/wines", json={
            "nome": "TEST_WithPhoto", "produttore": "TEST_PP",
            "annata": "2020", "foto_path": "cartadeivini/etichette/fake.jpg",
            "prezzo_carta": 10.0,
        }, timeout=30)
        assert r.status_code == 201, r.text
        data = _xlsx([
            ["Nome", "Produttore", "Annata", "Prezzo in carta"],
            ["TEST_WithPhoto", "TEST_PP", 2020, "15,00"],
        ])
        r = admin.post(f"{API}/import", files={"file": ("c.xlsx", data, XLSX_MIME)},
                       data={"mode": "update"}, timeout=30)
        assert r.status_code == 200
        assert r.json()["aggiornati"] == 1
        r = admin.get(f"{API}/wines", timeout=30)
        w = next(x for x in r.json() if x["nome"] == "TEST_WithPhoto")
        assert w["foto_path"] == "cartadeivini/etichette/fake.jpg"
        assert w["prezzo_carta"] == 15.0


# ---------- template + round-trip ----------
class TestTemplateAndRoundTrip:
    def test_template_is_valid_xlsx(self, admin):
        r = admin.get(f"{API}/import/template", timeout=30)
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        headers = [c.value for c in ws[1]]
        assert "Nome" in headers
        assert any("Produttore" in (h or "") for h in headers)

    def test_round_trip_xlsx_all_presenti(self, admin):
        # Export current wines; re-import; all should be "presente"
        r = admin.get(f"{API}/export/xlsx", timeout=30)
        assert r.status_code == 200
        r = admin.post(f"{API}/import/preview",
                       files={"file": ("export.xlsx", r.content, XLSX_MIME)}, timeout=30)
        assert r.status_code == 200
        t = r.json()["totali"]
        assert t["righe"] > 0
        assert t["nuovi"] == 0
        assert t["errori"] == 0
        assert t["presenti"] == t["righe"]

    def test_round_trip_csv_all_presenti(self, admin):
        r = admin.get(f"{API}/export/csv", timeout=30)
        assert r.status_code == 200
        r = admin.post(f"{API}/import/preview",
                       files={"file": ("export.csv", r.content, "text/csv")}, timeout=30)
        assert r.status_code == 200
        t = r.json()["totali"]
        assert t["righe"] > 0
        assert t["nuovi"] == 0
        assert t["presenti"] == t["righe"]
