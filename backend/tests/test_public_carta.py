"""Backend tests for Public Carta Digitale (iteration 5)."""
import os
import io
import csv
import uuid
import pytest
import requests
from openpyxl import load_workbook, Workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://carta-dei-vini.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"

ADMIN_EMAIL = "admin@cartadeivini.it"
ADMIN_PASSWORD = "Bordeaux-2026!"

STAFF_EMAIL = "TEST_pubstaff@cartadeivini.it"
STAFF_PASSWORD = "StaffPass-2026!"


@pytest.fixture(scope="module")
def admin():
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD}, timeout=30)
    assert r.status_code == 200, r.text
    yield s
    # ensure public is ON at end
    s.put(f"{API}/settings", json={"carta_pubblica_attiva": True}, timeout=30)


@pytest.fixture(scope="module")
def staff(admin):
    r = admin.post(f"{API}/users", json={
        "name": "TEST PubStaff", "email": STAFF_EMAIL,
        "password": STAFF_PASSWORD, "role": "staff",
    }, timeout=30)
    assert r.status_code in (201, 400), r.text
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json={"email": STAFF_EMAIL, "password": STAFF_PASSWORD}, timeout=30)
    assert r.status_code == 200
    yield s
    # cleanup user
    users = admin.get(f"{API}/users", timeout=30).json()
    for u in users:
        if u["email"] == STAFF_EMAIL:
            admin.delete(f"{API}/users/{u['id']}", timeout=30)


# shared test wines
TEST_WINES = {"visible": None, "hidden": None}


@pytest.fixture(scope="module", autouse=True)
def test_wines(admin):
    # visible wine
    r = admin.post(f"{API}/wines", json={
        "nome": "TEST_PUB_Visible", "produttore": "TEST_PUB_Cantina",
        "regione": "Toscana", "denominazione": "Chianti DOCG",
        "tipologia": "Rosso", "vitigno": "Sangiovese",
        "annata": "2020", "prezzo_acquisto": 10.0, "prezzo_carta": 30.0,
        "preferito": True, "visibile": True, "note": "Public test wine",
    }, timeout=30)
    assert r.status_code == 201, r.text
    TEST_WINES["visible"] = r.json()["id"]

    # hidden wine
    r = admin.post(f"{API}/wines", json={
        "nome": "TEST_PUB_Hidden", "produttore": "TEST_PUB_Cantina",
        "tipologia": "Bianco", "prezzo_acquisto": 5.0, "prezzo_carta": 20.0,
        "visibile": False,
    }, timeout=30)
    assert r.status_code == 201, r.text
    TEST_WINES["hidden"] = r.json()["id"]

    yield
    for wid in TEST_WINES.values():
        if wid:
            admin.delete(f"{API}/wines/{wid}", timeout=30)


# ---------- settings endpoint ----------
class TestSettings:
    def test_get_requires_auth(self):
        r = requests.get(f"{API}/settings", timeout=30)
        assert r.status_code == 401

    def test_get_as_admin(self, admin):
        r = admin.get(f"{API}/settings", timeout=30)
        assert r.status_code == 200
        assert "carta_pubblica_attiva" in r.json()

    def test_get_as_staff(self, staff):
        r = staff.get(f"{API}/settings", timeout=30)
        assert r.status_code == 200

    def test_put_staff_forbidden(self, staff):
        r = staff.put(f"{API}/settings", json={"carta_pubblica_attiva": False}, timeout=30)
        assert r.status_code == 403

    def test_put_admin_ok(self, admin):
        r = admin.put(f"{API}/settings", json={"carta_pubblica_attiva": True}, timeout=30)
        assert r.status_code == 200
        assert r.json()["carta_pubblica_attiva"] is True


# ---------- public carta endpoint ----------
class TestPublicCarta:
    def test_public_no_auth_needed(self, admin):
        # ensure public is on
        admin.put(f"{API}/settings", json={"carta_pubblica_attiva": True}, timeout=30)
        r = requests.get(f"{API}/public/carta", timeout=30)
        assert r.status_code == 200
        data = r.json()
        assert isinstance(data, list)
        ids = [w["id"] for w in data]
        assert TEST_WINES["visible"] in ids
        assert TEST_WINES["hidden"] not in ids

    def test_public_excludes_prezzo_acquisto_and_foto_path(self, admin):
        r = requests.get(f"{API}/public/carta", timeout=30)
        assert r.status_code == 200
        for w in r.json():
            assert "prezzo_acquisto" not in w, f"prezzo_acquisto leaked in {w.get('nome')}"
            assert "foto_path" not in w, f"foto_path leaked in {w.get('nome')}"
        # the visible test wine must have prezzo_carta
        v = next(w for w in r.json() if w["id"] == TEST_WINES["visible"])
        assert v["prezzo_carta"] == 30.0
        assert v["preferito"] is True
        assert "foto" in v  # boolean indicator

    def test_public_photo_hidden_wine_404(self, admin):
        r = requests.get(f"{API}/public/foto/{TEST_WINES['hidden']}", timeout=30)
        assert r.status_code == 404

    def test_public_photo_visible_no_foto_404(self, admin):
        # visible wine has no foto -> 404
        r = requests.get(f"{API}/public/foto/{TEST_WINES['visible']}", timeout=30)
        assert r.status_code == 404

    def test_public_photo_nonexistent_404(self):
        r = requests.get(f"{API}/public/foto/{uuid.uuid4()}", timeout=30)
        assert r.status_code == 404

    def test_toggle_off_returns_404_italian(self, admin):
        r = admin.put(f"{API}/settings", json={"carta_pubblica_attiva": False}, timeout=30)
        assert r.status_code == 200
        try:
            r = requests.get(f"{API}/public/carta", timeout=30)
            assert r.status_code == 404
            detail = r.json().get("detail", "")
            assert "carta digitale" in detail.lower()
            # also photo should 404
            r2 = requests.get(f"{API}/public/foto/{TEST_WINES['visible']}", timeout=30)
            assert r2.status_code == 404
        finally:
            admin.put(f"{API}/settings", json={"carta_pubblica_attiva": True}, timeout=30)


# ---------- visibility toggle flow ----------
class TestVisibility:
    def test_hide_then_show_wine(self, admin):
        wid = TEST_WINES["visible"]
        # hide via PUT
        admin.put(f"{API}/wines/{wid}", json={
            "nome": "TEST_PUB_Visible", "produttore": "TEST_PUB_Cantina",
            "visibile": False,
        }, timeout=30)
        r = requests.get(f"{API}/public/carta", timeout=30)
        ids = [w["id"] for w in r.json()]
        assert wid not in ids
        # unhide
        admin.put(f"{API}/wines/{wid}", json={
            "nome": "TEST_PUB_Visible", "produttore": "TEST_PUB_Cantina",
            "visibile": True,
        }, timeout=30)
        r = requests.get(f"{API}/public/carta", timeout=30)
        ids = [w["id"] for w in r.json()]
        assert wid in ids


# ---------- export/import with Visibile column ----------
class TestExportVisibile:
    def test_csv_has_visibile_column(self, admin):
        r = admin.get(f"{API}/export/csv", timeout=60)
        assert r.status_code == 200
        text = r.content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text), delimiter=";")
        headers = next(reader)
        assert "Visibile in carta digitale" in headers
        idx_vis = headers.index("Visibile in carta digitale")
        idx_nome = headers.index("Nome")
        rows = list(reader)
        hidden = next((row for row in rows if row[idx_nome] == "TEST_PUB_Hidden"), None)
        visible = next((row for row in rows if row[idx_nome] == "TEST_PUB_Visible"), None)
        assert hidden is not None and visible is not None
        assert hidden[idx_vis] == "No"
        assert visible[idx_vis] == "Sì"

    def test_xlsx_has_visibile_column(self, admin):
        r = admin.get(f"{API}/export/xlsx", timeout=60)
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        headers = [c.value for c in ws[1]]
        assert "Visibile in carta digitale" in headers

    def test_import_preview_maps_visibile(self, admin):
        # build a preview-only xlsx with Visibile column
        wb = Workbook()
        ws = wb.active
        headers = ["Nome", "Produttore / Cantina", "Visibile in carta digitale"]
        ws.append(headers)
        ws.append(["TEST_PUB_Visible", "TEST_PUB_Cantina", "No"])  # existing -> different vis
        ws.append(["TEST_PUB_Hidden", "TEST_PUB_Cantina", ""])  # empty -> should treat visible
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        r = admin.post(
            f"{API}/import/preview",
            files={"file": ("preview.xlsx", buf.getvalue(),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            timeout=60,
        )
        assert r.status_code == 200, r.text
        data = r.json()
        # Row-level results should exist for both rows, no error about Visibile mapping
        rows = data.get("righe") or data.get("rows") or []
        assert len(rows) >= 2
        # Scan for any error strings mentioning column mapping issue
        joined = str(data).lower()
        assert "visibile" not in joined or "non riconosciut" not in joined
